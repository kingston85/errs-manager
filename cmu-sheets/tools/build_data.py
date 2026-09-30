#!/usr/bin/env python3
"""
Build the CMU Sheets dataset from the original workbook.

Reads  "CMU_DATABASE Original.xlsx"  (repo root, or path given as argv[1])
Writes cmu-sheets/data/cmu-base.json   — all registers, lists, fees, settings
       cmu-sheets/data/macro-code.txt  — the original VBA code (Macro Setup sheet)

Run:  python3 tools/build_data.py
"""
import datetime as dt
import json
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "CMU_DATABASE Original.xlsx")
OUT_DIR = os.path.join(os.path.dirname(HERE), "data")


def cell_value(v):
    """Normalise a workbook value for JSON."""
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        if v.time() == dt.time(0, 0):
            return v.date().isoformat()
        return v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return int(v)
        return v
    if isinstance(v, str):
        s = v.strip()
        return s if s else None
    return v


def read_rows(ws, header_row, first_col, last_col, max_rows=None, stop_after_blank=15):
    """Read a table: headers at header_row, data below until long blank run."""
    headers = []
    for c in range(first_col, last_col + 1):
        h = cell_value(ws.cell(row=header_row, column=c).value)
        headers.append(h or "")
    rows = []
    blank = 0
    r = header_row + 1
    while r <= ws.max_row:
        vals = [cell_value(ws.cell(row=r, column=c).value) for c in range(first_col, last_col + 1)]
        if all(v is None for v in vals):
            blank += 1
            if blank >= stop_after_blank:
                break
        else:
            blank = 0
            rows.append(vals)
        r += 1
        if max_rows and len(rows) >= max_rows:
            break
    # trim trailing all-None cells per row (keep at least the id)
    trimmed = []
    for vals in rows:
        while len(vals) > 1 and vals[-1] is None:
            vals.pop()
        trimmed.append(vals)
    return headers, trimmed


def pad(rows, ncols):
    out = []
    for vals in rows:
        vals = list(vals) + [None] * (ncols - len(vals))
        out.append(vals)
    return out


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)

    data = {
        "meta": {
            "title": "EPA LIBERIA – ERRS | CHEMICAL MANAGEMENT UNIT (CMU) DATABASE 2026",
            "source": os.path.basename(SRC),
            "built": dt.date.today().isoformat(),
            "codeVersion": "2026.09.30-UI",
        },
        "sheets": {},
        "lists": {},
        "fees": {"services": [], "billTypes": [], "footnote": ""},
        "printSettings": {},
        "users": [],
    }

    # ------------------------------------------------------------------ registers
    REGISTERS = {
        "companies": ("Companies", 4, 1, 10),
        "licences": ("Licences & Certificates", 4, 1, 17),
        "licenceChemicals": ("Licence Chemicals", 4, 1, 9),
        "licenceVehicles": ("Licence Vehicles", 4, 1, 6),
        "licenceEffluent": ("Licence Effluent", 4, 1, 12),
        "clearances": ("Chemical Clearances", 4, 1, 17),
        "bills": ("Bills & Invoices", 4, 1, 22),
        "payments": ("Payments & Receipts", 4, 1, 12),
    }
    for key, (sheet, hrow, c1, c2) in REGISTERS.items():
        ws = wb[sheet]
        headers, rows = read_rows(ws, hrow, c1, c2)
        rows = [r for r in rows if not (r and isinstance(r[0], str) and r[0].upper() == "TOTAL")]
        rows = pad(rows, c2 - c1 + 1)
        data["sheets"][key] = {
            "name": sheet,
            "title": cell_value(ws["A1"].value) or sheet,
            "subtitle": cell_value(ws["A2"].value) or "",
            "headers": headers,
            "rows": rows,
        }

    # Notes & Issues holds two tables: issues (A4:G) and correspondence (A30:I)
    ws = wb["Issues_ws_tmp"] if "Issues_ws_tmp" in wb.sheetnames else wb["Notes & Issues"]
    headers, rows = read_rows(ws, 4, 1, 7)
    rows = [r for r in rows if r and isinstance(r[0], str) and r[0].startswith("ISS")]
    data["sheets"]["issues"] = {
        "name": "Notes & Issues", "title": cell_value(ws["A1"].value) or "Notes & Issues",
        "subtitle": cell_value(ws["A2"].value) or "", "headers": headers, "rows": pad(rows, 7),
    }
    headers, rows = read_rows(ws, 30, 1, 9)
    data["sheets"]["correspondence"] = {
        "name": "Correspondence", "title": "ERRS – CHEMICAL MANAGEMENT UNIT (CMU) | CORRESPONDENCE LOG 2026",
        "subtitle": "Reminder letters, violation notices, shutdown notices (from the Notes & Issues sheet)",
        "headers": headers, "rows": pad(rows, 9),
    }

    # ------------------------------------------------------------------ source data
    ws = wb["SRC BMMC Log"]
    headers, rows = read_rows(ws, 3, 1, 4)
    data["sheets"]["srcBmmc"] = {
        "name": "SRC BMMC Log", "title": cell_value(ws["A1"].value) or "SRC BMMC Log", "subtitle": "Original BMMC clearance log – kept for reference",
        "headers": headers, "rows": pad(rows, 4),
    }
    ws = wb["SRC CMU Receipts"]
    headers, rows = read_rows(ws, 2, 1, 5)
    data["sheets"]["srcReceipts"] = {
        "name": "SRC CMU Receipts", "title": cell_value(ws["A1"].value) or "SRC CMU Receipts", "subtitle": "Original CMU receipts log – kept for reference",
        "headers": headers, "rows": pad(rows, 5),
    }

    # ------------------------------------------------------------------ lists
    ws = wb["Lists"]
    def col_list(col, start=5, stop_at_blank=40):
        vals, blank = [], 0
        r = start
        while r <= ws.max_row:
            v = cell_value(ws.cell(row=r, column=col).value)
            if v is None:
                blank += 1
                if blank >= stop_at_blank:
                    break
            else:
                blank = 0
                vals.append(v)
            r += 1
        return vals

    data["lists"] = {
        "licenceCategory": col_list(1),
        "licenceType": col_list(3),
        "units": col_list(5),
        "chemicals": col_list(7),
        "sectors": col_list(9),
        "servicePurpose": col_list(11),
        "recordGroup": col_list(13),
        "parameters": col_list(39),
        "effluentUnits": col_list(41),
        "testingSchedules": col_list(43),
        "vehicleTypes": ["Heavy Duty Truck", "Tanker Truck", "Pick-up", "Van", "Other"],
    }

    # signatories: T (label), U (name), V (title)
    sigs = []
    for r in range(5, 40):
        label = cell_value(ws.cell(row=r, column=20).value)
        name = cell_value(ws.cell(row=r, column=21).value)
        title = cell_value(ws.cell(row=r, column=22).value)
        if label or name:
            sigs.append({"label": label or name, "name": name or label, "title": title or ""})
    data["lists"]["signatories"] = sigs

    # signatory offices: AF..AI
    offices = []
    for r in range(5, 40):
        sig = cell_value(ws.cell(row=r, column=32).value)
        office = cell_value(ws.cell(row=r, column=33).value)
        contact = cell_value(ws.cell(row=r, column=34).value)
        prefix = cell_value(ws.cell(row=r, column=35).value)
        if sig:
            offices.append({"signatory": sig, "officeLine": office or "", "contactLine": contact or "", "refPrefix": prefix or ""})
    data["lists"]["offices"] = offices

    # print settings: O (label) / P (value), rows 5..30
    for r in range(5, 31):
        k = cell_value(ws.cell(row=r, column=15).value)
        v = cell_value(ws.cell(row=r, column=16).value)
        if k and v is not None:
            data["printSettings"][k] = v

    # ------------------------------------------------------------------ fee schedule
    ws = wb["Fee Schedule"]
    for r in range(6, 17):
        svc = cell_value(ws.cell(row=r, column=1).value)
        if svc:
            data["fees"]["services"].append({
                "service": svc,
                "group": cell_value(ws.cell(row=r, column=2).value) or "",
                "fee": cell_value(ws.cell(row=r, column=3).value),
                "notes": cell_value(ws.cell(row=r, column=4).value) or "",
            })
    for r in range(20, 27):
        t = cell_value(ws.cell(row=r, column=1).value)
        if t:
            data["fees"]["billTypes"].append({
                "type": t,
                "group": cell_value(ws.cell(row=r, column=2).value) or "",
                "opening": cell_value(ws.cell(row=r, column=3).value) or "",
                "extra": cell_value(ws.cell(row=r, column=4).value) or "",
                "billSentence": cell_value(ws.cell(row=r, column=5).value) or "",
                "collection": cell_value(ws.cell(row=r, column=6).value) or "",
                "showFeeTable": (cell_value(ws.cell(row=r, column=7).value) or "No") == "Yes",
            })
    data["fees"]["footnote"] = cell_value(ws.cell(row=28, column=1).value) or ""

    # ------------------------------------------------------------------ users (names/roles only – fresh credentials are seeded by the server)
    ws = wb["Users"]
    for r in range(5, 60):
        u = cell_value(ws.cell(row=r, column=1).value)
        if u:
            data["users"].append({
                "username": u,
                "fullName": cell_value(ws.cell(row=r, column=2).value) or u,
                "role": cell_value(ws.cell(row=r, column=3).value) or "Officer",
                "active": (cell_value(ws.cell(row=r, column=4).value) or "Yes") == "Yes",
            })

    # ------------------------------------------------------------------ VBA code (Macro Setup, column A from row 16)
    ws = wb["Macro Setup"]
    lines = []
    for r in range(16, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        lines.append("" if v is None else str(v))
    macro = "\n".join(lines).rstrip() + "\n"
    header = (
        "' ====================================================================\n"
        "' ORIGINAL VBA / MACRO CODE from \"CMU_DATABASE Original.xlsx\" › Macro Setup\n"
        "' (kept for reference – the CMU Sheets web app replaces all of this code)\n"
        "' ====================================================================\n\n"
    )
    os.makedirs(OUT_DIR, exist_ok=True)
    with open(os.path.join(OUT_DIR, "macro-code.txt"), "w") as f:
        f.write(header + macro)
    with open(os.path.join(OUT_DIR, "cmu-base.json"), "w") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))

    print("Wrote", os.path.join(OUT_DIR, "cmu-base.json"))
    for k, v in data["sheets"].items():
        print(f"  {k:18s} {v['name'][:30]:32s} rows={len(v['rows'])}")
    print("  lists:", {k: len(v) for k, v in data["lists"].items()})
    print("  fees: services=%d billTypes=%d" % (len(data["fees"]["services"]), len(data["fees"]["billTypes"])))
    print("  macro code lines:", macro.count("\n"))


if __name__ == "__main__":
    main()
