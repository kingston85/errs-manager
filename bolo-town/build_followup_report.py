"""
Render the Bolo Town follow-up investigative assessment report (content in
report_content.py) to DOCX and PDF, following the layout of the initial
investigation report (cover page, table of contents, blue-banner section
headings, Times body text, page-numbered footer, photo appendix).

    python3 build_followup_report.py
"""
import re
from pathlib import Path

# ---- DOCX -----------------------------------------------------------------
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ---- PDF ------------------------------------------------------------------
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (BaseDocTemplate, Frame, Image, KeepTogether, ListFlowable,
                                ListItem, NextPageTemplate, PageBreak, PageTemplate,
                                Paragraph, Spacer, Table, TableStyle)
from reportlab.platypus.tableofcontents import TableOfContents

from report_content import BLOCKS, COVER, ABBREVIATIONS

OUT_DIR = Path(__file__).resolve().parent
BASENAME = "Bolo Town Follow-up Investigative Assessment Report (Oct 13 2026) - DRAFT"

BANNER = "4F81BD"      # heading banner blue (as in the initial report)
HDR_FILL = "404040"    # table header fill
ROW_SHADE = "F2F2F2"
PLACEHOLDER_GREY = "7F7F7F"

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _runs(text):
    """Split '**bold**' markup into [(text, is_bold), ...]."""
    out, pos = [], 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False))
        out.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False))
    return out


# ===========================================================================
# DOCX renderer
# ===========================================================================
def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _para_shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _field(paragraph, instr, placeholder=""):
    r = paragraph.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); b.set(qn("w:dirty"), "true")
    r._r.append(b)
    r = paragraph.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    r._r.append(it)
    r = paragraph.add_run()
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); r._r.append(s)
    paragraph.add_run(placeholder)
    r = paragraph.add_run()
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._r.append(e)


def _add_text(paragraph, text, size=None, bold=None, italic=None, color=None):
    for chunk, is_b in _runs(text):
        run = paragraph.add_run(chunk)
        if size: run.font.size = Pt(size)
        if bold or is_b: run.bold = True
        if italic: run.italic = True
        if color: run.font.color.rgb = RGBColor.from_string(color)
    return paragraph


def _docx_table(doc, spec):
    header, rows = spec["header"], spec["rows"]
    ncols = len(header[0]); nh = len(header)
    t = doc.add_table(rows=nh + len(rows), cols=ncols)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fs = spec["font"]
    # merges
    for (c0, r0), (c1, r1) in spec.get("spans", []):
        t.cell(r0, c0).merge(t.cell(r1, c1))
    for (c0, r0), (c1, r1) in spec.get("body_spans", []):
        t.cell(r0, c0).merge(t.cell(r1, c1))
    # header text
    done = set()
    for r, hrow in enumerate(header):
        for c, txt in enumerate(hrow):
            cell = t.cell(r, c)
            if id(cell._tc) in done or not txt:
                continue
            done.add(id(cell._tc))
            cell.text = ""
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_text(p, txt, size=fs, bold=True, color="FFFFFF")
    for r in range(nh):
        for c in range(ncols):
            _shade(t.cell(r, c), HDR_FILL)
    # body
    for i, row in enumerate(rows):
        for c, txt in enumerate(row):
            cell = t.cell(nh + i, c)
            if id(cell._tc) in done:
                continue
            done.add(id(cell._tc))
            cell.text = ""
            p = cell.paragraphs[0]
            if c in spec["center_cols"] or txt == "\u2014":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_text(p, txt, size=fs)
            if i in spec["shade_rows"]:
                _shade(cell, ROW_SHADE)
    for row in t.rows:
        for c, w in enumerate(spec["widths_in"]):
            row.cells[c].width = Inches(w)
    # repeat header rows
    if spec.get("repeat_header"):
        for r in range(nh):
            trPr = t.rows[r]._tr.get_or_add_trPr()
            el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)
    return t


def build_docx(path: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(0.75)
    sec.top_margin = sec.bottom_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for lvl, size in ((1, 12), (2, 11), (3, 11)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Times New Roman"; st.font.size = Pt(size); st.font.bold = True
        st.font.color.rgb = RGBColor.from_string("FFFFFF" if lvl == 1 else "1F3864")
        st.element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        st.element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")

    # footer page numbers
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ").font.size = Pt(9)
    _field(fp, "PAGE", "1")

    # ---- cover ---------------------------------------------------------
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run().add_picture(str(COVER["logo"]), width=Inches(4.2))
    p = doc.add_paragraph(); r = p.add_run(COVER["tagline"]); r.italic = True; r.font.size = Pt(9)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(COVER["title"]); r.bold = True; r.font.size = Pt(15)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(COVER["subtitle"]); r.italic = True; r.font.size = Pt(11.5)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(COVER["photo"]), width=Inches(5.2))
    doc.add_paragraph(); doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(COVER["prepared_by"]); r.bold = True; r.font.size = Pt(12)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(COVER["date"]); r.bold = True; r.font.size = Pt(12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(COVER["status"]); r.italic = True; r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string("C00000")
    doc.add_page_break()

    # ---- table of contents ---------------------------------------------
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Table of Contents"); r.bold = True; r.font.size = Pt(12)
    # TOC field whose pre-update result is a static list of the headings, so the
    # page reads correctly even before Word refreshes the field (F9 adds page numbers).
    p = doc.add_paragraph()
    r = p.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); b.set(qn("w:dirty"), "true"); r._r.append(b)
    r = p.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = 'TOC \\o "1-2" \\h \\z \\u'
    r._r.append(it)
    r = p.add_run(); sp = OxmlElement("w:fldChar"); sp.set(qn("w:fldCharType"), "separate"); r._r.append(sp)
    first = True
    for blk in BLOCKS:
        if blk[0] not in ("h1", "h2"):
            continue
        if not first:
            p = doc.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(2)
        if blk[0] == "h2":
            p.paragraph_format.left_indent = Inches(0.3)
        run = p.add_run(blk[1]); run.font.size = Pt(10.5 if blk[0] == "h1" else 10)
        run.bold = blk[0] == "h1"
    r = p.add_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._r.append(e)
    note = doc.add_paragraph(); note.paragraph_format.space_before = Pt(8)
    _add_text(note, "Page numbers: right-click the list above and choose \u201cUpdate Field\u201d (or press F9).",
              size=8.5, italic=True, color=PLACEHOLDER_GREY)
    # ask Word to refresh fields on open
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); settings.append(uf)
    doc.add_page_break()

    # ---- list of abbreviations (front matter) ---------------------------
    h = doc.add_paragraph(); h.paragraph_format.space_after = Pt(8)
    _add_text(h, "LIST OF ABBREVIATIONS", size=12, bold=True, color="FFFFFF"); _para_shade(h, BANNER)
    at = doc.add_table(rows=1, cols=2); at.style = "Table Grid"; at.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, txt in enumerate(("Abbreviation", "Meaning")):
        cell = at.cell(0, c); cell.text = ""; _shade(cell, HDR_FILL)
        _add_text(cell.paragraphs[0], txt, size=10, bold=True, color="FFFFFF")
    for abbr, meaning in ABBREVIATIONS:
        cells = at.add_row().cells
        cells[0].text = ""; _add_text(cells[0].paragraphs[0], abbr, size=10, bold=True)
        cells[1].text = ""; _add_text(cells[1].paragraphs[0], meaning, size=10)
    for row in at.rows:
        row.cells[0].width = Inches(1.5); row.cells[1].width = Inches(5.2)
    doc.add_page_break()

    # ---- body ----------------------------------------------------------
    for blk in BLOCKS:
        kind = blk[0]
        if kind == "h1":
            h = doc.add_heading(level=1); h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            _add_text(h, blk[1].upper()); _para_shade(h, BANNER)
        elif kind == "h2":
            h = doc.add_heading(level=2); _add_text(h, blk[1].upper())
            h.paragraph_format.space_before = Pt(10); h.paragraph_format.space_after = Pt(4)
        elif kind == "h3":
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
            _add_text(p, blk[1], bold=True)
        elif kind == "p":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.15
            _add_text(p, blk[1])
        elif kind in ("bullets", "numbered"):
            style = "List Bullet" if kind == "bullets" else "List Number"
            for item in blk[1]:
                p = doc.add_paragraph(style=style); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_after = Pt(3); _add_text(p, item)
        elif kind == "placeholder":
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
            _add_text(p, f"[To be added: {blk[1]}]", italic=True, color=PLACEHOLDER_GREY)
        elif kind == "caption":
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2); p.paragraph_format.keep_with_next = True
            _add_text(p, blk[1], size=9.5, bold=True)
        elif kind == "table":
            _docx_table(doc, blk[1]); doc.add_paragraph()
        elif kind == "photos":
            items = blk[1]
            t = doc.add_table(rows=(len(items) + 1) // 2, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, (fp_, capt) in enumerate(items):
                cell = t.cell(i // 2, i % 2); cell.text = ""
                cp = cell.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.add_run().add_picture(str(fp_), height=Inches(1.8))
                cp2 = cell.add_paragraph(); cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _add_text(cp2, capt, size=8.5, italic=True)
            doc.add_paragraph()
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "spacer":
            doc.add_paragraph().paragraph_format.space_after = Pt(blk[1])
    doc.save(path)


# ===========================================================================
# PDF renderer
# ===========================================================================
def _rl(text):
    """Escape for reportlab Paragraph and convert **bold**."""
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return BOLD_RE.sub(r"<b>\1</b>", text)


class ReportDoc(BaseDocTemplate):
    def __init__(self, filename, **kw):
        super().__init__(filename, **kw)
        frame = Frame(self.leftMargin, self.bottomMargin, self.width, self.height, id="f")
        self.addPageTemplates([
            PageTemplate(id="cover", frames=[frame]),
            PageTemplate(id="body", frames=[frame], onPageEnd=self._footer),
        ])
        self._body_start = None

    def _pageno(self):
        return self.page - (self._body_start or self.page) + 1

    def _footer(self, canv, doc):
        if self._body_start is None or self._pageno() < 1:
            return
        canv.saveState(); canv.setFont("Times-Roman", 9)
        canv.drawCentredString(A4[0] / 2, 0.45 * inch, f"Page {self._pageno()}")
        canv.restoreState()

    def afterFlowable(self, fl):
        if hasattr(fl, "_toc_level"):
            if self._body_start is None:
                self._body_start = self.page
            key = f"h{id(fl)}"
            self.canv.bookmarkPage(key)
            self.notify("TOCEntry", (fl._toc_level, fl._toc_text, self._pageno(), key))


def build_pdf(path: Path):
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Times-Roman", fontSize=11,
                          leading=15, alignment=TA_JUSTIFY, spaceAfter=6)
    h1 = ParagraphStyle("h1", parent=body, fontName="Times-Bold", fontSize=12, leading=15,
                        textColor=colors.white, backColor=colors.HexColor("#" + BANNER),
                        borderPadding=(3, 6, 3, 6), spaceBefore=14, spaceAfter=10, alignment=0)
    h2 = ParagraphStyle("h2", parent=body, fontName="Times-Bold", fontSize=11, leading=14,
                        textColor=colors.HexColor("#1F3864"), spaceBefore=10, spaceAfter=4, alignment=0)
    h3 = ParagraphStyle("h3", parent=body, fontName="Times-Bold", spaceBefore=8, spaceAfter=3, alignment=0)
    ph = ParagraphStyle("ph", parent=body, fontName="Times-Italic", textColor=colors.HexColor("#" + PLACEHOLDER_GREY))
    cap = ParagraphStyle("cap", parent=body, fontName="Times-Bold", fontSize=9.5, leading=12,
                         spaceBefore=6, spaceAfter=3, alignment=0, keepWithNext=True)
    cell = ParagraphStyle("cell", parent=body, fontSize=8, leading=10, alignment=0, spaceAfter=0)
    cell_c = ParagraphStyle("cell_c", parent=cell, alignment=TA_CENTER)
    cell_h = ParagraphStyle("cell_h", parent=cell_c, fontName="Times-Bold", textColor=colors.white)
    cover_title = ParagraphStyle("ct", parent=body, fontName="Times-Bold", fontSize=15, leading=20, alignment=TA_CENTER)
    cover_sub = ParagraphStyle("cs", parent=body, fontName="Times-Italic", fontSize=11.5, leading=15, alignment=TA_CENTER)
    cover_b = ParagraphStyle("cb", parent=body, fontName="Times-Bold", fontSize=12, leading=16, alignment=TA_CENTER)
    cover_status = ParagraphStyle("cst", parent=body, fontName="Times-Italic", fontSize=9.5,
                                  textColor=colors.HexColor("#C00000"), alignment=TA_CENTER)
    toc_h = ParagraphStyle("toch", parent=body, fontName="Times-Bold", fontSize=12, alignment=TA_CENTER, spaceAfter=10)
    photo_cap = ParagraphStyle("pc", parent=cell_c, fontName="Times-Italic", fontSize=8.5, leading=10.5)

    doc = ReportDoc(str(path), pagesize=A4, leftMargin=0.75 * inch, rightMargin=0.75 * inch,
                    topMargin=0.8 * inch, bottomMargin=0.8 * inch, title=BASENAME)
    story = []

    # ---- cover ---------------------------------------------------------
    logo = Image(str(COVER["logo"])); logo.drawWidth, logo.drawHeight = 4.2 * inch, 4.2 * inch * logo.imageHeight / logo.imageWidth
    logo.hAlign = "LEFT"; story.append(logo)
    story.append(Paragraph(f"<i>{COVER['tagline']}</i>", ParagraphStyle("tag", parent=body, fontSize=9, alignment=0)))
    story.append(Spacer(1, 24))
    story.append(Paragraph(COVER["title"], cover_title)); story.append(Spacer(1, 8))
    story.append(Paragraph(COVER["subtitle"], cover_sub)); story.append(Spacer(1, 22))
    photo = Image(str(COVER["photo"])); photo.drawWidth = 5.2 * inch
    photo.drawHeight = 5.2 * inch * photo.imageHeight / photo.imageWidth; story.append(photo)
    story.append(Spacer(1, 40))
    story.append(Paragraph(COVER["prepared_by"], cover_b)); story.append(Spacer(1, 18))
    story.append(Paragraph(COVER["date"], cover_b)); story.append(Spacer(1, 6))
    story.append(Paragraph(COVER["status"], cover_status))
    story.append(NextPageTemplate("body")); story.append(PageBreak())

    # ---- TOC -----------------------------------------------------------
    story.append(Paragraph("Table of Contents", toc_h))
    toc = TableOfContents()
    toc.levelStyles = [
        ParagraphStyle("toc1", parent=body, fontName="Times-Bold", fontSize=10.5, leading=13.5, leftIndent=0, alignment=0, spaceAfter=0),
        ParagraphStyle("toc2", parent=body, fontName="Times-Roman", fontSize=10, leading=12.5, leftIndent=18, alignment=0, spaceAfter=0),
    ]
    toc.dotsMinLevel = 0
    story.append(toc); story.append(PageBreak())

    # ---- list of abbreviations (front matter, not in TOC) --------------
    story.append(Paragraph("LIST OF ABBREVIATIONS", h1))
    abbr_cell = ParagraphStyle("ab", parent=cell, fontSize=10, leading=13)
    abbr_b = ParagraphStyle("abb", parent=abbr_cell, fontName="Times-Bold")
    data = [[Paragraph("Abbreviation", ParagraphStyle("abh", parent=abbr_b, textColor=colors.white)),
             Paragraph("Meaning", ParagraphStyle("abh2", parent=abbr_b, textColor=colors.white))]]
    data += [[Paragraph(_rl(a), abbr_b), Paragraph(_rl(m), abbr_cell)] for a, m in ABBREVIATIONS]
    at = Table(data, colWidths=[1.5 * inch, 5.2 * inch], repeatRows=1)
    at.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.6, colors.black),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + HDR_FILL)),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    story.append(at); story.append(PageBreak())

    # ---- body ----------------------------------------------------------
    def heading(text, level):
        p = Paragraph(_rl(text.upper()), h1 if level == 0 else h2)
        p._toc_level, p._toc_text = level, text
        return p

    def table(spec):
        header, rows = spec["header"], spec["rows"]
        nh = len(header)
        data = []
        for hrow in header:
            data.append([Paragraph(_rl(x), cell_h) if x else "" for x in hrow])
        for row in rows:
            data.append([Paragraph(_rl(x), cell_c if (c in spec["center_cols"] or x == "\u2014") else cell)
                         if x else "" for c, x in enumerate(row)])
        fs = spec["font"]
        for r in data:
            for pg in r:
                if isinstance(pg, Paragraph):
                    pg.style = ParagraphStyle("tmp", parent=pg.style, fontSize=fs, leading=fs + 2.5)
        t = Table(data, colWidths=[w * inch for w in spec["widths_in"]],
                  repeatRows=nh if spec.get("repeat_header") else 0)
        ts = TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
            ("BACKGROUND", (0, 0), (-1, nh - 1), colors.HexColor("#" + HDR_FILL)),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ])
        for (c0, r0), (c1, r1) in spec.get("spans", []) + spec.get("body_spans", []):
            ts.add("SPAN", (c0, r0), (c1, r1))
        for i in spec["shade_rows"]:
            ts.add("BACKGROUND", (0, nh + i), (-1, nh + i), colors.HexColor("#" + ROW_SHADE))
        t.setStyle(ts); t.hAlign = "CENTER"
        return t

    pending_caption = None
    for blk in BLOCKS:
        kind = blk[0]
        if kind == "h1":
            story.append(heading(blk[1], 0))
        elif kind == "h2":
            story.append(heading(blk[1], 1))
        elif kind == "h3":
            story.append(Paragraph(_rl(blk[1]), h3))
        elif kind == "p":
            story.append(Paragraph(_rl(blk[1]), body))
        elif kind in ("bullets", "numbered"):
            items = [ListItem(Paragraph(_rl(x), body), leftIndent=18) for x in blk[1]]
            if kind == "bullets":
                story.append(ListFlowable(items, bulletType="bullet", start="\u2022", leftIndent=18, bulletFontSize=9))
            else:
                story.append(ListFlowable(items, bulletType="1", bulletFormat="%s.", leftIndent=18,
                                          bulletFontName="Times-Roman", bulletFontSize=11))
        elif kind == "placeholder":
            story.append(Paragraph(_rl(f"[To be added: {blk[1]}]"), ph))
        elif kind == "caption":
            pending_caption = Paragraph(_rl(blk[1]), cap)
        elif kind == "table":
            t = table(blk[1])
            if pending_caption is not None:
                story.append(KeepTogether([pending_caption, Spacer(1, 2)]))
                pending_caption = None
            story.append(t); story.append(Spacer(1, 8))
        elif kind == "photos":
            cells, pw, phh = [], 3.3 * inch, 1.8 * inch
            for fp_, capt in blk[1]:
                im = Image(str(fp_)); im.drawHeight = phh; im.drawWidth = phh * im.imageWidth / im.imageHeight
                cells.append([im, Spacer(1, 2), Paragraph(_rl(capt), photo_cap)])
            rows = [cells[i:i + 2] for i in range(0, len(cells), 2)]
            for r in rows:
                while len(r) < 2:
                    r.append("")
            g = Table(rows, colWidths=[pw + 0.15 * inch] * 2)
            g.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
            story.append(g)
        elif kind == "pagebreak":
            story.append(PageBreak())
        elif kind == "spacer":
            story.append(Spacer(1, blk[1]))

    doc.multiBuild(story)


if __name__ == "__main__":
    docx_path = OUT_DIR / f"{BASENAME}.docx"
    pdf_path = OUT_DIR / f"{BASENAME}.pdf"
    build_docx(docx_path); print("wrote", docx_path)
    build_pdf(pdf_path); print("wrote", pdf_path)
